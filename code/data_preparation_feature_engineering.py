from __future__ import annotations

import re
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

ONET_DIR = ROOT / "data" / "raw" / "onet"
BLS_FILE = ROOT / "data" / "raw" / "bls" / "all_data_M_2025.xlsx"
OUT = ROOT / "data"

OUT.mkdir(parents=True, exist_ok=True)
OUT.mkdir(exist_ok=True)


def snake(text: str) -> str:
    text = re.sub(r'[^0-9A-Za-z]+', '_', str(text).strip().lower())
    return re.sub(r'_+', '_', text).strip('_')


def normalize_cols(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [snake(c) for c in df.columns]
    return df


def read_onet(name: str) -> pd.DataFrame:
    csv_path = ROOT / 'csv_onet' / name.replace('.xlsx', '.csv')
    df = pd.read_csv(csv_path) if csv_path.exists() else pd.read_excel(ONET_DIR / name)
    df = normalize_cols(df)
    if 'o_net_soc_code' in df:
        df['onet_soc_code'] = df['o_net_soc_code'].astype('string').str.strip()
        df = df.drop(columns=['o_net_soc_code'])
        df['soc_code'] = df['onet_soc_code'].str.extract(r'^(\d{2}-\d{4})', expand=False)
    return df


def read_bls_national() -> pd.DataFrame:
    wb = load_workbook(BLS_FILE, read_only=True, data_only=True)
    ws = wb['All May 2025 data']
    rows = ws.iter_rows(values_only=True)
    header = [snake(x) for x in next(rows)]
    keep = []
    idx = {c: i for i, c in enumerate(header)}
    started = False
    for row in rows:
        area_type = str(row[idx['area_type']])
        if area_type == '1':
            started = True
            # U.S. national, cross-industry, all ownerships, detailed occupations only.
            if (str(row[idx['area']]) == '99'
                    and str(row[idx['naics']]) == '000000'
                    and str(row[idx['own_code']]) == '1235'
                    and str(row[idx['o_group']]).lower() == 'detailed'):
                keep.append(row)
        elif started:
            # National rows are the first contiguous block in the OEWS workbook.
            break
    wb.close()
    df = pd.DataFrame(keep, columns=header)
    df = df.rename(columns={'occ_code': 'soc_code', 'occ_title': 'bls_occ_title'})
    df['soc_code'] = df['soc_code'].astype('string').str.strip()
    numeric = ['tot_emp','emp_prse','h_mean','a_mean','mean_prse','h_pct10','h_pct25',
               'h_median','h_pct75','h_pct90','a_pct10','a_pct25','a_median','a_pct75','a_pct90']
    for c in numeric:
        df[c] = pd.to_numeric(df[c], errors='coerce')
    return df


def pivot_ratings(file_name: str, prefix: str, scales=('IM','LV')) -> pd.DataFrame:
    df = read_onet(file_name)
    df = df[df['scale_id'].isin(scales)].copy()
    df = df[df['recommend_suppress'].fillna('N').eq('N')] if 'recommend_suppress' in df else df
    df['feature'] = prefix + '_' + df['element_name'].map(snake) + '_' + df['scale_id'].str.lower()
    wide = df.pivot_table(index='onet_soc_code', columns='feature', values='data_value', aggfunc='mean')
    return wide.reset_index()


def add_domain_summary(master: pd.DataFrame, prefix: str) -> pd.DataFrame:
    cols = [c for c in master.columns if c.startswith(prefix + '_') and c.endswith('_im')]
    if cols:
        master[f'{prefix}_importance_mean'] = master[cols].mean(axis=1)
        master[f'{prefix}_importance_max'] = master[cols].max(axis=1)
    lv = [c for c in master.columns if c.startswith(prefix + '_') and c.endswith('_lv')]
    if lv:
        master[f'{prefix}_level_mean'] = master[lv].mean(axis=1)
        master[f'{prefix}_level_max'] = master[lv].max(axis=1)
    return master

# Base occupation table.
occ = read_onet('Occupation Data.xlsx')
occ = occ.rename(columns={'title':'occupation_title', 'description':'occupation_description'})
occ = occ[['onet_soc_code','soc_code','occupation_title','occupation_description']]
occ = occ.drop_duplicates('onet_soc_code')

# One-row-per-occupation O*NET dimensions.
parts = [
    pivot_ratings('Essential Skills.xlsx', 'essential_skill'),
    pivot_ratings('Transferable Skills.xlsx', 'transferable_skill'),
    pivot_ratings('Knowledge.xlsx', 'knowledge'),
    pivot_ratings('Abilities.xlsx', 'ability'),
    pivot_ratings('Work Activities.xlsx', 'work_activity'),
]

# Interests (RIASEC).
interests = read_onet('Career Interest Types.xlsx')
interests = interests[interests['scale_id'].eq('OI')].copy()
interests['feature'] = 'interest_' + interests['element_name'].map(snake)
interests_w = interests.pivot_table(index='onet_soc_code', columns='feature', values='data_value', aggfunc='mean').reset_index()
parts.append(interests_w)

# Work styles: use impact scale; distinctiveness rank is retained separately only if needed.
styles = read_onet('Work Styles.xlsx')
styles = styles[styles['scale_id'].eq('WI')].copy()
styles['feature'] = 'work_style_' + styles['element_name'].map(snake) + '_impact'
styles_w = styles.pivot_table(index='onet_soc_code', columns='feature', values='data_value', aggfunc='mean').reset_index()
parts.append(styles_w)

# Job zone.
jobz = read_onet('Job Zones.xlsx')[['onet_soc_code','job_zone']].drop_duplicates('onet_soc_code')
parts.append(jobz)

# Education: expected category, modal category, and bachelor's-or-higher response share.
edu = read_onet('Education.xlsx')
edu['category'] = pd.to_numeric(edu['category'], errors='coerce')
edu['data_value'] = pd.to_numeric(edu['data_value'], errors='coerce')
edu = edu[edu['scale_id'].eq('RL')]
edu['weighted'] = edu['category'] * edu['data_value']
edu_summary = edu.groupby('onet_soc_code').agg(
    education_expected_category=('weighted','sum'),
    education_response_total=('data_value','sum')
).reset_index()
edu_summary['education_expected_category'] = edu_summary['education_expected_category'] / edu_summary['education_response_total'].replace(0,np.nan)
modal_idx = edu.groupby('onet_soc_code')['data_value'].idxmax()
modal = edu.loc[modal_idx, ['onet_soc_code','category']].rename(columns={'category':'education_modal_category'})
bach = edu[edu['category'].ge(6)].groupby('onet_soc_code', as_index=False)['data_value'].sum().rename(columns={'data_value':'education_bachelors_or_higher_pct'})
edu_summary = edu_summary.merge(modal,on='onet_soc_code',how='left').merge(bach,on='onet_soc_code',how='left')
parts.append(edu_summary)

master = occ.copy()
for p in parts:
    master = master.merge(p, on='onet_soc_code', how='left', validate='one_to_one')

for prefix in ['essential_skill','transferable_skill','knowledge','ability','work_activity']:
    master = add_domain_summary(master, prefix)

# Dominant RIASEC interest and spread.
interest_cols = [c for c in master.columns if c.startswith('interest_')]
if interest_cols:
    master['dominant_interest'] = master[interest_cols].idxmax(axis=1).str.replace('interest_','',regex=False)
    master['interest_score_max'] = master[interest_cols].max(axis=1)
    master['interest_score_range'] = master[interest_cols].max(axis=1) - master[interest_cols].min(axis=1)

# Merge with BLS national detailed occupation estimates.
bls = read_bls_national()
bls_cols = ['soc_code','bls_occ_title','tot_emp','emp_prse','h_mean','a_mean','mean_prse',
            'h_pct10','h_pct25','h_median','h_pct75','h_pct90','a_pct10','a_pct25',
            'a_median','a_pct75','a_pct90','annual','hourly']
bls = bls[bls_cols].drop_duplicates('soc_code')
master = master.merge(bls, on='soc_code', how='left', validate='many_to_one', indicator=True)
master['bls_match'] = master['_merge'].eq('both')
master = master.drop(columns=['_merge'])

# Labor-market engineered features.
master['employment_log1p'] = np.log1p(master['tot_emp'])
master['annual_wage_spread_90_10'] = master['a_pct90'] - master['a_pct10']
master['annual_wage_iqr'] = master['a_pct75'] - master['a_pct25']
master['median_to_mean_wage_ratio'] = master['a_median'] / master['a_mean'].replace(0,np.nan)
master['wage_data_available'] = master['a_median'].notna()

# Percentile ranks are interpretable, scale-free opportunity features.
for src, dest in [('a_median','median_wage_percentile'),('tot_emp','employment_percentile')]:
    master[dest] = master[src].rank(pct=True, na_option='keep')
master['labor_market_opportunity_score'] = master[['median_wage_percentile','employment_percentile']].mean(axis=1)

# Quality reports.
quality = []
def q(check, value, action): quality.append({'check':check,'result':value,'action':action})
q('O*NET occupation rows', len(occ), 'Retained one row per full O*NET-SOC code')
q('Unique O*NET-SOC codes', occ['onet_soc_code'].nunique(), 'Expected to equal occupation rows')
q('Duplicate full O*NET-SOC codes', int(occ['onet_soc_code'].duplicated().sum()), 'Removed exact identifier duplicates')
q('BLS national detailed occupation rows', len(bls), 'Filtered U.S., cross-industry, all ownership, detailed occupations')
q('Unique BLS SOC codes', bls['soc_code'].nunique(), 'Used as many-to-one merge key')
q('Matched O*NET occupations', int(master['bls_match'].sum()), 'Retained BLS labor-market measures')
q('Unmatched O*NET occupations', int((~master['bls_match']).sum()), 'Retained for review; no silent deletion')
q('BLS match rate', f"{master['bls_match'].mean():.1%}", 'Documented merge coverage')
q('Negative employment values', int((master['tot_emp'] < 0).sum()), 'None expected')
q('Negative annual wage values', int((master['a_median'] < 0).sum()), 'None expected')
q('Final rows', len(master), 'One row per O*NET occupation')
q('Final columns', master.shape[1], 'Includes identifiers, O*NET features, BLS measures, engineered features')
quality_df = pd.DataFrame(quality)

missing = master.isna().sum().rename('missing_count').to_frame()
missing['missing_percent'] = missing['missing_count'] / len(master) * 100
missing = missing.sort_values('missing_percent', ascending=False).reset_index(names='variable')

unmatched = master.loc[~master['bls_match'], ['onet_soc_code','soc_code','occupation_title']]

# Feature dictionary.
source_map = []
for c in master.columns:
    if c in ['onet_soc_code','soc_code','occupation_title','occupation_description']:
        source='O*NET Occupation Data'
    elif c.startswith(('essential_skill_','transferable_skill_','knowledge_','ability_','work_activity_','interest_','work_style_')):
        source='O*NET content model table'
    elif c.startswith('education_') or c=='job_zone':
        source='O*NET education/job zone'
    elif c in bls_cols or c in ['bls_match']:
        source='BLS May 2025 OEWS'
    else:
        source='Engineered'
    modeling = 'Identifier/display only' if c in ['onet_soc_code','soc_code','occupation_title','occupation_description','bls_occ_title'] else 'Candidate model/recommendation feature'
    source_map.append({'variable':c,'source':source,'dtype':str(master[c].dtype),'role':modeling})
feature_dict = pd.DataFrame(source_map)

# Outputs.
master.to_csv(OUT / 'career_compass_master.csv', index=False)
quality_df.to_csv(OUT / 'data_quality_report.csv', index=False)
missing.to_csv(OUT / 'missing_values_report.csv', index=False)
unmatched.to_csv(OUT / 'unmatched_onet_occupations.csv', index=False)
feature_dict.to_csv(OUT / 'feature_dictionary.csv', index=False)
bls.to_csv(OUT / 'bls_national_detailed_clean.csv', index=False)

summary = f"""# Data Preparation and Feature Engineering Results\n\n- O*NET release: 30.3\n- BLS source: May 2025 OEWS All Data\n- O*NET occupations: {len(occ):,}\n- BLS national detailed occupations: {len(bls):,}\n- Final master rows: {len(master):,}\n- Final master columns: {master.shape[1]:,}\n- O*NET occupations matched to BLS: {int(master['bls_match'].sum()):,} ({master['bls_match'].mean():.1%})\n- Unmatched O*NET occupations: {int((~master['bls_match']).sum()):,}\n\nThe final dataset preserves one row per full O*NET-SOC occupation. The six-digit SOC code is used only as the many-to-one bridge to national BLS estimates.\n"""
(OUT / 'results_summary.md').write_text(summary, encoding='utf-8')
print(summary)
print(quality_df.to_string(index=False))
